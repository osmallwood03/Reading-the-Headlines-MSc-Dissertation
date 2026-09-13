"""
Recover HM Treasury's aggregation rule rather than assuming one.

    python validate_aggregation.py --panel <panel.csv> --workbook <hmt.xlsx>

The panel gives individual forecasts; E_P needs one number per publication per
event. HMT published its own average for 2010-10 to 2023-04, so candidate rules
are scored against that benchmark and the one that reproduces it is adopted.

The naive answer is badly wrong in a way that matters here. HMT reprints
forecasts contributors have not updated, so a mean over everyone listed is
dragged towards stale beliefs, and dragged furthest when beliefs are moving
fastest.

Two things govern how to read the output. The benchmark is printed to one
decimal place, so a perfect rule still scores about 0.025 on mean absolute
error; that is the floor, not a good result. And error correlated with the
state of the world biases the estimate while uncorrelated error only adds
noise, so the per-year table and the crisis/quiet split carry the argument,
not the headline mean.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import sys
from pathlib import Path

import numpy as np
import pandas as pd

from build_panel import load_panel

# Rounding floor of the benchmark: HMT prints 1 d.p., so E|error| >= 0.025
# even for a perfect rule.
ROUNDING_FLOOR = 0.025
WITHIN_ROUNDING = 0.051   # a hit = indistinguishable from the printed figure

# The OBR is listed in the panel but is not an independent forecaster in the
# sense the rest of the panel means: it is the government's official forecast,
# produced at fiscal events rather than on the monthly submission cycle, and
# it is policy-conditioned. Whether HMT's published average includes it is an
# empirical question, tested below rather than assumed.
OFFICIAL_FORECASTERS = ["OBR"]


# Benchmark: the published averages in the HMT workbook

def load_workbook_average(xlsx_path: str | Path,
                          sheets: tuple[str, ...] = ("RPI", "CPI")
                          ) -> pd.DataFrame:
    """
    Read the published average from the 'Database of Average Forecasts'.

    Layout: rows are publication months, columns are fixed events labelled
    'YYYY Q4' (a Q4-on-Q4 annual inflation rate -- a point in time, not a
    calendar-year average). Cells are the published average, 1 d.p.

    A handful of cells carry text flags instead of numbers ('NO PUBLICATION',
    'MISSING'); these are skipped, not coerced to zero.
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    out = []
    for sheet in sheets:
        rows = list(wb[sheet].iter_rows(values_only=True))
        hdr = [str(h).replace(" ", "").replace("Q4", "") for h in rows[0]]
        for r in rows[1:]:
            d = r[0]
            if not isinstance(d, _dt.datetime):
                continue
            for h, v in zip(hdr[1:], r[1:]):
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    out.append((pd.Timestamp(d), f"{sheet} (Q4)", int(h), v))
    return pd.DataFrame(out, columns=["pub", "variable", "event_year",
                                      "published"])


# Candidate aggregation rules

def _rules() -> dict[str, dict]:
    """
    Each rule = a row filter plus a central-tendency statistic.

    `max_age` is the staleness window in months: how old a forecast may be and
    still enter the average. age 0 = submitted in the publication month (HMT
    marks these with an asterisk).
    """
    r: dict[str, dict] = {}
    r["mean, everyone listed"] = dict(max_age=None, stat="mean", drop_obr=False)
    for a in (0, 1, 2, 3, 6):
        r[f"mean, age <= {a}m"] = dict(max_age=a, stat="mean", drop_obr=False)
    r["mean, age <= 2m, ex-OBR"] = dict(max_age=2, stat="mean", drop_obr=True)
    r["mean, everyone, ex-OBR"] = dict(max_age=None, stat="mean", drop_obr=True)
    r["median, age <= 2m"] = dict(max_age=2, stat="median", drop_obr=False)
    r["median, everyone listed"] = dict(max_age=None, stat="median",
                                        drop_obr=False)
    return r


def apply_rule(ind: pd.DataFrame, max_age: int | None, stat: str,
               drop_obr: bool) -> pd.DataFrame:
    """Aggregate the panel to one value per (pub, variable, event_year)."""
    d = ind
    if drop_obr:
        d = d[~d["forecaster"].isin(OFFICIAL_FORECASTERS)]
    if max_age is not None:
        # Age filters are meaningless where the submission column is known bad
        # (see UNRELIABLE_SUBMISSION_VINTAGES in build_panel), so those months
        # are excluded from the comparison rather than scored on junk dates.
        d = d[d["subm_reliable"] & (d["age_m"] <= max_age)]
    g = d.groupby(["pub", "variable", "event_year"])["value"]
    out = (g.mean() if stat == "mean" else g.median()).rename("fitted")
    return pd.concat([out, g.size().rename("n")], axis=1).reset_index()


# Scoring

def score(ind: pd.DataFrame, bench: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for name, cfg in _rules().items():
        fit = apply_rule(ind, **cfg)
        m = bench.merge(fit, on=["pub", "variable", "event_year"], how="inner")
        m = m[m["fitted"].notna()]
        if not len(m):
            continue
        e = (m["published"] - m["fitted"]).abs()
        crisis = m["pub"].dt.year.isin([2021, 2022, 2023])
        rows.append({
            "rule": name,
            "n": len(m),
            "mean|err|": e.mean(),
            "med|err|": e.median(),
            "hit rate": (e <= WITHIN_ROUNDING).mean(),
            "max|err|": e.max(),
            "err 2021-23": e[crisis].mean(),
            "err other": e[~crisis].mean(),
            "median n": m["n"].median(),
        })
    return pd.DataFrame(rows).sort_values("mean|err|")


def regime_check(ind: pd.DataFrame, bench: pd.DataFrame, cfg: dict) -> tuple:
    """
    Is the error of the chosen rule correlated with the state of the world?

    This is the test that decides whether residual error is benign noise or
    attenuating bias. `published` is the professional forecast itself, so it
    proxies the inflation regime directly.
    """
    fit = apply_rule(ind, **cfg)
    m = bench.merge(fit, on=["pub", "variable", "event_year"], how="inner")
    m = m[m["fitted"].notna()].copy()
    m["err"] = m["published"] - m["fitted"]
    r_signed = np.corrcoef(m["published"], m["err"])[0, 1]
    r_abs = np.corrcoef(m["published"], m["err"].abs())[0, 1]
    return r_signed, r_abs, m


# Report

def _h(t: str) -> None:
    print("\n" + "=" * 82)
    print(t)
    print("=" * 82)


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--panel", default="panel_forecasters_june2026_in_.csv")
    p.add_argument("--workbook",
                   default="Database_of_Average_Forecasts_for_the_UK_Economy.xlsx")
    a = p.parse_args()
    for f in (a.panel, a.workbook):
        if not Path(f).exists():
            print(f"ERROR: cannot find {f}", file=sys.stderr)
            return 1

    ind, _, _ = load_panel(a.panel)
    q4 = ind[ind["variable"].str.contains(r"\(Q4\)")].copy()

    # Guard: the undated rows must never reach E^P silently.
    n_undated = int(q4["subm"].isna().sum())
    assert n_undated == 0, (
        f"{n_undated} Q4 rows have no submission date. Every age filter would "
        "drop them silently (NaN comparisons are False). Decide explicitly "
        "before proceeding."
    )

    bench = load_workbook_average(a.workbook)
    ov = bench.merge(q4[["pub", "variable", "event_year"]].drop_duplicates(),
                     on=["pub", "variable", "event_year"], how="inner")

    _h("OVERLAP USED AS THE BENCHMARK")
    print(f"  published cells matched to panel cells : {len(ov):,}")
    print(f"  publication months                     : {ov['pub'].nunique()}")
    print(f"  range                                  : "
          f"{ov['pub'].min():%Y-%m} .. {ov['pub'].max():%Y-%m}")
    print(f"  guard passed: no undated rows in the Q4 series")

    _h("CANDIDATE RULES, SCORED AGAINST THE PUBLISHED AVERAGE")
    print(f"  Rounding floor of the benchmark: mean|err| cannot beat "
          f"{ROUNDING_FLOOR:.3f}")
    print(f"  'hit rate' = share within +/-{WITHIN_ROUNDING} of the printed "
          f"figure (i.e. indistinguishable from it)\n")
    tab = score(ind, bench)
    print(tab.to_string(index=False, float_format=lambda x: f"{x:6.3f}"))

    best_name = tab.iloc[0]["rule"]
    best_cfg = _rules()[best_name]
    _h(f"REGIME-DEPENDENCE OF THE LEADING RULE: {best_name}")
    r_signed, r_abs, m = regime_check(ind, bench, best_cfg)
    print("  Error that is correlated with the inflation regime biases the")
    print("  estimated effect; error that is not merely adds noise.\n")
    print(f"  corr(published, signed error) = {r_signed:+.3f}")
    print(f"  corr(published, |error|)      = {r_abs:+.3f}")

    _h("PER-YEAR ERROR OF THE LEADING RULE  (read this, not the headline)")
    m["yr"] = m["pub"].dt.year
    t = m.groupby("yr").agg(n=("err", "size"),
                            mean_abs=("err", lambda s: s.abs().mean()),
                            mean_signed=("err", "mean"),
                            worst=("err", lambda s: s.abs().max()))
    print(f"  {'year':<6}{'n':>5}{'mean|err|':>11}{'mean err':>10}{'worst':>8}")
    for y, r in t.iterrows():
        flag = "  <-- above rounding floor" if r["mean_abs"] > 0.06 else ""
        print(f"  {y:<6}{int(r['n']):>5}{r['mean_abs']:>11.3f}"
              f"{r['mean_signed']:>10.3f}{r['worst']:>8.3f}{flag}")

    _h("WHAT THE NAIVE RULE WOULD HAVE COST YOU  (RPI, 2022)")
    naive = apply_rule(ind, max_age=None, stat="mean", drop_obr=False)
    best = apply_rule(ind, **best_cfg)
    cmp = (bench[(bench["variable"] == "RPI (Q4)")
                 & (bench["pub"].dt.year == 2022)]
           .merge(naive, on=["pub", "variable", "event_year"])
           .merge(best, on=["pub", "variable", "event_year"],
                  suffixes=("_naive", "_best")))
    print(f"  {'pub':<9}{'event':>7}{'published':>11}{'naive mean':>12}"
          f"{'chosen rule':>13}{'naive gap':>11}")
    for _, r in cmp.sort_values(["pub", "event_year"]).iterrows():
        print(f"  {r['pub']:%Y-%m}   {r['event_year']:>6}"
              f"{r['published']:>11.2f}{r['fitted_naive']:>12.2f}"
              f"{r['fitted_best']:>13.2f}"
              f"{r['published'] - r['fitted_naive']:>11.2f}")
    d = (cmp["published"] - cmp["fitted_naive"])
    print(f"\n  Mean understatement from the naive rule in 2022: "
          f"{d.mean():.2f} pp (worst {d.max():.2f} pp)")
    print("  For scale: the household-professional gap you are explaining")
    print("  peaks at roughly 2 pp.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
